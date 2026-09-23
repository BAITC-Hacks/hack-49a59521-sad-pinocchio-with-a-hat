from io import BytesIO
from pathlib import Path
from threading import Lock
from xml.sax.saxutils import escape

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from app.config import Settings
from app.errors import ProviderError
from app.models import Meeting

FONT_LOCK = Lock()


def font_name(settings: Settings) -> str:
    candidates = [
        settings.pdf_font_path,
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    path = next((Path(p) for p in candidates if p and Path(p).is_file()), None)
    if path is None:
        raise ProviderError("Не найден Unicode-шрифт: задайте PDF_FONT_PATH (DejaVu Sans TTF)")
    with FONT_LOCK:
        name = "MeetingUnicode_" + str(abs(hash(str(path.resolve()))))
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    return name


def timestamp(seconds: float) -> str:
    return f"{int(seconds) // 3600:02}:{int(seconds) // 60 % 60:02}:{seconds % 60:04.1f}"


def speaker_names(meeting: Meeting) -> dict:
    return {s.id: s.name or s.id for s in meeting.speakers}


def task_rows(meeting: Meeting) -> list[list[str]]:
    names = speaker_names(meeting)
    rows = [["Поручение", "Ответственный", "Срок", "Статус"]]
    for task in meeting.tasks:
        detail = task.title
        if task.description and task.description != task.title:
            detail += "\n" + task.description
        detail += f"\nПриоритет: {task.priority}\nИсточник: {', '.join(task.source_segment_ids)}"
        rows.append(
            [
                detail,
                task.assignee_name or names.get(task.assignee, "Не назначен"),
                str(task.deadline or "Не указан"),
                task.status,
            ]
        )
    return rows


def transcript_lines(meeting: Meeting):
    names = speaker_names(meeting)
    for seg in meeting.transcript.segments if meeting.transcript else []:
        yield (
            f"[{timestamp(seg.start)} - {timestamp(seg.end)}] "
            f"{names.get(seg.speaker_id, seg.speaker_id)} ({seg.id}): {seg.text}"
        )


def export_pdf(meeting: Meeting, settings: Settings) -> bytes:
    font = font_name(settings)
    normal = ParagraphStyle(
        "body",
        fontName=font,
        fontSize=9,
        leading=13,
        spaceAfter=6,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    heading = ParagraphStyle(
        "heading",
        parent=normal,
        fontSize=13,
        leading=17,
        spaceBefore=12,
        spaceAfter=8,
        keepWithNext=True,
    )
    title = ParagraphStyle("title", parent=heading, fontSize=19, leading=24)

    def p(text, style=normal):
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), style)

    content = [
        p(meeting.title, title),
        p(f"Дата совещания: {meeting.date}"),
        p(f"ID: {meeting.meeting_id}"),
        p("Краткое содержание", heading),
        p(meeting.summary or "Нет данных"),
    ]
    for label, items in [
        ("Решения", meeting.decisions),
        ("Нерешённые вопросы", meeting.unresolved_questions),
        ("Примечания к обработке", meeting.warnings),
    ]:
        if items:
            content.append(p(label, heading))
            content.extend(p(f"{i + 1}. {text}") for i, text in enumerate(items))
    content.append(p("Поручения", heading))
    if meeting.tasks:
        table = LongTable(
            [[p(cell) for cell in row] for row in task_rows(meeting)],
            colWidths=[83 * mm, 36 * mm, 27 * mm, 28 * mm],
            repeatRows=1,
            splitInRow=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#c3cbd5")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 6),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        content.append(table)
    else:
        content.append(p("Поручения не выявлены"))
    content.extend([Spacer(1, 5 * mm), p("Полный транскрипт", heading)])
    content.extend(p(line) for line in transcript_lines(meeting))
    output = BytesIO()

    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.drawRightString(192 * mm, 10 * mm, f"Страница {doc.page}")

    SimpleDocTemplate(
        output,
        pagesize=(210 * mm, 297 * mm),
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=17 * mm,
        bottomMargin=18 * mm,
        title=meeting.title,
    ).build(content, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_docx(meeting: Meeting, settings: Settings) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(1.8)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.name = "DejaVu Sans"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_paragraph(meeting.title, "Title")
    doc.add_paragraph(f"Дата совещания: {meeting.date}\nID: {meeting.meeting_id}")
    doc.add_heading("Краткое содержание", 1)
    doc.add_paragraph(meeting.summary or "Нет данных")
    for label, items in [
        ("Решения", meeting.decisions),
        ("Нерешённые вопросы", meeting.unresolved_questions),
        ("Примечания к обработке", meeting.warnings),
    ]:
        if items:
            doc.add_heading(label, 1)
            for item in items:
                doc.add_paragraph(item, "List Bullet")
    doc.add_heading("Поручения", 1)
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, [8.3, 3.6, 2.7, 2.8]):
        column.width = Cm(width)
    for values in task_rows(meeting):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, values, [8.3, 3.6, 2.7, 2.8]):
            cell.width = Cm(width)
            cell.text = value
    header = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.add_heading("Полный транскрипт", 1)
    for line in transcript_lines(meeting):
        doc.add_paragraph(line)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def prepare_exports(meeting: Meeting, settings: Settings, directory: Path):
    directory.mkdir(parents=True, exist_ok=True)
    for extension, render in [("pdf", export_pdf), ("docx", export_docx)]:
        temporary = directory / f"minutes.{extension}.tmp"
        temporary.write_bytes(render(meeting, settings))
        temporary.replace(directory / f"minutes.{extension}")
