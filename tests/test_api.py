from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfReader

from app.errors import ProviderError


def upload(client, name="meeting.mp3", content=b"mock-audio", mime="audio/mpeg"):
    return client.post(
        "/api/meetings/upload",
        files={"file": (name, content, mime)},
        data={"meeting_date": "2026-09-23", "title": "Тестовое совещание"},
    )


def test_upload_process_and_persistence(client):
    response = upload(client)
    assert response.status_code == 201
    data = response.json()
    mid = data["meeting_id"]
    assert data["status"] == "uploaded"
    assert client.get(f"/api/meetings/{mid}/status").json()["status"] == "uploaded"
    assert client.get(f"/api/meetings/{mid}/export/pdf").status_code == 409
    assert client.get(f"/api/meetings/{mid}/transcript").status_code == 409
    response = client.post(f"/api/meetings/{mid}/process")
    assert response.status_code == 202
    assert response.json()["status"] == "processing"
    assert client.get(f"/api/meetings/{mid}/status").json()["status"] == "completed"
    meeting = client.get(f"/api/meetings/{mid}").json()
    assert len(meeting["tasks"]) == 5
    assert meeting["duration"] == 63
    assert meeting["date"] == "2026-09-23"
    assert client.post(f"/api/meetings/{mid}/process").status_code == 409
    assert (client.app.state.storage.results / mid / "minutes.pdf").is_file()
    assert (client.app.state.storage.results / mid / "minutes.docx").is_file()


@pytest.mark.parametrize(
    "name,content,mime,status",
    [
        ("file.exe", b"x", "application/octet-stream", 415),
        ("file.mp3", b"x", "text/plain", 415),
        ("file.wav", b"", "audio/wav", 400),
        ("file.wav", b"x" * (1024 * 1024 + 1), "audio/wav", 413),
    ],
    ids=["extension", "mime", "empty", "oversize"],
)
def test_invalid_upload(client, name, content, mime, status):
    assert upload(client, name, content, mime).status_code == status
    assert not list(client.app.state.storage.uploads.iterdir())


def test_missing_file_and_safe_filename(client):
    assert client.post("/api/meetings/upload").status_code == 422
    result = upload(client, "../../CON.mp3").json()
    assert result["filename"] == "upload_CON.mp3"
    path = client.app.state.storage.uploads / result["meeting_id"] / result["filename"]
    assert path.read_bytes() == b"mock-audio"


def test_demo_and_transcript(client, demo):
    assert demo["status"] == "completed"
    assert demo["language"] == "mixed"
    assert len(demo["speakers"]) >= 3
    assert demo["decisions"] and demo["unresolved_questions"] and demo["summary"]
    tasks = demo["tasks"]
    assert any(t["deadline"] == "2026-09-30" for t in tasks)
    assert any(t["deadline"] == "2026-09-24" for t in tasks)
    assert any(t["deadline"] is None for t in tasks)
    assert any(t["assignee"] is None for t in tasks)
    result = client.get(f"/api/meetings/{demo['meeting_id']}/transcript")
    assert result.json() == demo["transcript"]


def test_patch_task_speaker_and_fresh_exports(client, demo):
    base = f"/api/meetings/{demo['meeting_id']}"
    response = client.patch(base + "/speakers/speaker_2", json={"name": "Әлия"})
    assert response.json() == {"id": "speaker_2", "name": "Әлия"}
    response = client.patch(
        base + "/tasks/task-1",
        json={"status": "done", "priority": "high", "deadline": None, "title": "Жаңа есеп"},
    )
    assert response.status_code == 200
    assert response.json()["assignee_name"] == "Әлия"
    assert response.json()["deadline"] is None
    pdf = client.get(base + "/export/pdf")
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages)
    assert "Әлия" in pdf_text and "Жаңа есеп" in pdf_text and "done" in pdf_text
    assert "Полный транскрипт" in pdf_text and "00:00:00.0" in pdf_text
    docx = client.get(base + "/export/docx")
    assert docx.status_code == 200
    doc = Document(BytesIO(docx.content))
    table_text = " ".join(c.text for r in doc.tables[0].rows for c in r.cells)
    assert "Әлия" in table_text and "Жаңа есеп" in table_text
    assert "2026-09-30" not in doc.tables[0].rows[1].cells[2].text
    assert any("00:00:00.0" in p.text for p in doc.paragraphs)
    assert (
        client.patch(base + "/tasks/task-1", json={"assignee": None}).json()["assignee_name"]
        is None
    )


@pytest.mark.parametrize(
    "body",
    [
        {"status": "invalid"},
        {"title": None},
        {"priority": None},
        {"confidence": 1},
        {"deadline": "tomorrow"},
        {"assignee": "speaker_999"},
    ],
)
def test_invalid_patch(client, demo, body):
    assert (
        client.patch(f"/api/meetings/{demo['meeting_id']}/tasks/task-1", json=body).status_code
        == 422
    )


def test_not_found(client, demo):
    assert client.get(f"/api/meetings/{uuid4()}").status_code == 404
    assert client.get("/api/meetings/not-a-uuid").status_code == 422
    base = f"/api/meetings/{demo['meeting_id']}"
    assert client.patch(base + "/tasks/missing", json={"status": "done"}).status_code == 404
    assert client.patch(base + "/speakers/missing", json={"name": "A"}).status_code == 404


def test_failure_and_retry(client, monkeypatch):
    def fail(*args):
        raise ProviderError("Тестовая ошибка STT")

    mid = upload(client).json()["meeting_id"]
    with monkeypatch.context() as patch:
        patch.setattr("app.stt.MockSpeechToTextProvider.transcribe", fail)
        assert client.post(f"/api/meetings/{mid}/process").status_code == 202
    status = client.get(f"/api/meetings/{mid}/status").json()
    assert status["status"] == "failed" and status["error"] == "Тестовая ошибка STT"
    client.post(f"/api/meetings/{mid}/process")
    assert client.get(f"/api/meetings/{mid}/status").json()["status"] == "completed"


def test_demo_export_failure_is_not_success(client, monkeypatch):
    def fail(*args):
        raise ProviderError("Экспорт недоступен")

    monkeypatch.setattr("app.pipeline.prepare_exports", fail)
    response = client.post("/api/demo/meeting")
    assert response.status_code == 500
    assert response.json()["status"] == "failed"
